"""
Export service: turns a lit review draft (markdown with ## headers and
[paper_id] citations) into a downloadable .docx file — the "client-ready
demo" deliverable from the proposal's Week 4 goal.

This does a lightweight markdown -> docx conversion by hand rather than
pulling in a full markdown-parsing dependency, since the lit review agent
only ever produces a small, predictable set of markdown patterns (##
headers and plain paragraphs).
"""

import os
import uuid
from docx import Document
from docx.shared import Pt

EXPORT_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "generated_exports")


def _ensure_export_dir():
    os.makedirs(EXPORT_DIR, exist_ok=True)


def export_lit_review_to_docx(title: str, lit_review_markdown: str) -> str:
    """
    Converts the lit review markdown into a .docx file and returns the
    absolute file path it was saved to.
    """
    _ensure_export_dir()

    doc = Document()

    title_heading = doc.add_heading(title, level=0)

    for line in lit_review_markdown.split("\n"):
        stripped = line.strip()

        if not stripped:
            continue

        if stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=1)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("- "):
            doc.add_paragraph(stripped[2:], style="List Bullet")
        else:
            paragraph = doc.add_paragraph(stripped)
            paragraph.paragraph_format.space_after = Pt(8)

    safe_name = "".join(c if c.isalnum() or c in "-_ " else "" for c in title)[:50]
    filename = f"{safe_name.strip().replace(' ', '_')}_{uuid.uuid4().hex[:8]}.docx"
    filepath = os.path.join(EXPORT_DIR, filename)

    doc.save(filepath)
    return os.path.abspath(filepath)
